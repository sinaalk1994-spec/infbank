from io import BytesIO
from tempfile import TemporaryDirectory
from django.test import TestCase, Client, override_settings
from django.contrib.auth.models import User, Group
from django.core.exceptions import ValidationError, PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from .models import *
from .services import *

class WorkflowTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin=User.objects.create_superuser('admin','', 'Test-only-password-928!')
        cls.clerk=User.objects.create_user('clerk',password='Test-only-password-929!')
        cls.accountant=User.objects.create_user('accountant',password='Test-only-password-930!')
        cls.viewer=User.objects.create_user('viewer',password='Test-only-password-931!')
        for u,role in [(cls.clerk,'clerk'),(cls.accountant,'accountant'),(cls.viewer,'viewer')]:
            group,_=Group.objects.get_or_create(name=role); u.groups.add(group)
        cls.grade=Grade.objects.create(name='درجه یک')
        cls.member=Member.objects.create(name='عضو آزمایشی',national_id='0012345678')
        cls.lic=License.objects.create(member=cls.member,number='001234',business_name='واحد آزمایشی',grade=cls.grade)
        Settings.current()
    def due(self,year=1405,amount=5000000):
        return Due.objects.create(license=self.lic,year=year,amount=amount,basis='آزمایش')
    def pay(self,amount,due=None):
        return receive(self.accountant,self.lic,amount,'card','TEST','1405/01/15',due_id=due.pk if due else None)
    def test_partial_payment_discount_and_settlement(self):
        due=self.due()
        discount(self.admin,due,0,'مصوبه رئیس',percent=10)
        self.pay(2000000)
        self.assertEqual(due.balance,2500000); self.assertFalse(due.settled)
        self.pay(2500000); self.assertTrue(due.settled)
    def test_past_debt_paid_first_and_credit_preserved(self):
        opening=self.due(None,100); past=self.due(1403,200); current=self.due(1405,300)
        receipt=self.pay(650)
        self.assertTrue(all(d.settled for d in [opening,past,current]))
        self.assertEqual(receipt.available,50)
    def test_selected_year_does_not_clear_past(self):
        past=self.due(1403,100); current=self.due(1405,300)
        self.pay(300,current)
        self.assertTrue(current.settled); self.assertEqual(past.balance,100)
    def test_reopen_keeps_receipt_and_audits_approver(self):
        due=self.due(amount=1000); discount(self.admin,due,100,'تخفیف مجاز'); receipt=self.pay(900)
        reopen(self.accountant,self.admin,due,'ثبت اشتباه سال عضویت')
        self.assertEqual(due.balance,1000); self.assertEqual(receipt.available,900)
        self.assertEqual(Receipt.objects.count(),1); self.assertEqual(Allocation.objects.count(),1)
        entry=Audit.objects.first(); self.assertEqual(entry.after['approver'],'admin'); self.assertEqual(entry.actor,self.accountant)
        receive(self.accountant,self.lic,0,'','','',receipt=receipt)
        self.assertEqual(due.balance,100); self.assertEqual(receipt.available,0)
    def test_non_admin_cannot_approve_reopen(self):
        due=self.due(amount=100); self.pay(100)
        with self.assertRaises(PermissionDenied): reopen(self.accountant,self.clerk,due,'اصلاح اشتباه')
        self.assertTrue(due.settled)
    def test_wrong_admin_password_does_not_change_money(self):
        due=self.due(amount=100); self.pay(100); self.client.force_login(self.accountant)
        response=self.client.post(reverse('reopen',args=[self.lic.pk,due.pk]),{'admin_username':'admin','admin_password':'wrong','reason':'اصلاح اشتباه'})
        self.assertEqual(response.status_code,302); self.assertTrue(due.settled)
    def test_valid_admin_credentials_allow_reopen_via_http(self):
        due=self.due(amount=100); self.pay(100); self.client.force_login(self.accountant)
        response=self.client.post(reverse('reopen',args=[self.lic.pk,due.pk]),{'admin_username':'admin','admin_password':'Test-only-password-928!','reason':'اصلاح اشتباه'})
        self.assertEqual(response.status_code,302); self.assertEqual(due.balance,100)
    def test_reopen_cannot_be_replayed(self):
        due=self.due(amount=100); self.pay(100); reopen(self.accountant,self.admin,due,'اصلاح اشتباه')
        with self.assertRaises(ValidationError): reopen(self.accountant,self.admin,due,'تلاش دوباره')
    def test_no_over_discount(self):
        due=self.due(amount=100); self.pay(70)
        with self.assertRaises(ValidationError): discount(self.admin,due,31,'تخفیف')
        self.assertEqual(due.balance,30)
    def test_tariff_is_snapshot_and_billing_idempotent(self):
        tariff=Tariff.objects.create(year=1405,grade=self.grade,amount=500)
        self.assertEqual(bill_year(self.accountant,1405),1)
        tariff.amount=1000; tariff.save()
        self.assertEqual(bill_year(self.accountant,1405),0)
        self.assertEqual(Due.objects.get().amount,500)
    def test_billing_all_or_nothing_for_missing_tariff(self):
        with self.assertRaises(ValidationError): bill_year(self.admin,1405)
        self.assertEqual(Due.objects.count(),0)
    def test_role_boundaries(self):
        for user,path in [(self.clerk,'payment'),(self.viewer,'member_edit')]:
            self.client.force_login(user)
            self.assertEqual(self.client.post(reverse(path,args=[self.lic.pk]),{}).status_code,403)
        self.client.force_login(self.accountant)
        self.assertEqual(self.client.get(reverse('settings')).status_code,403)
        self.client.force_login(self.clerk)
        self.assertEqual(self.client.get(reverse('export')).status_code,403)
    def test_login_lockout(self):
        for i in range(5): self.client.post('/login/',{'username':'admin','password':'wrong'})
        response=self.client.post('/login/',{'username':'admin','password':'Test-only-password-928!'})
        self.assertEqual(response.status_code,200); self.assertNotIn('_auth_user_id',self.client.session)
    def test_anonymous_and_csrf(self):
        self.assertEqual(self.client.get('/').status_code,302)
        client=Client(enforce_csrf_checks=True); client.force_login(self.admin)
        self.assertEqual(client.post(reverse('payment',args=[self.lic.pk]),{}).status_code,403)
    def test_all_pages_render_and_only_latest_year_in_list(self):
        self.due(1403,200); current=self.due(1405,300); self.pay(300,current)
        self.client.force_login(self.admin)
        for url in ['/','/members/','/members/new/',f'/members/{self.lic.pk}/',f'/members/{self.lic.pk}/edit/','/finance/','/archive/','/excel/','/settings/','/users/','/audit/']:
            with self.subTest(url=url): self.assertEqual(self.client.get(url).status_code,200)
        response=self.client.get('/members/')
        self.assertContains(response,'تسویه ۱۴۰۵'); self.assertNotContains(response,'۱۴۰۳')
        self.assertContains(self.client.get(f'/members/{self.lic.pk}/'),'۱۴۰۳')
    def test_excel_export_preserves_zero_and_neutralizes_formula(self):
        self.member.name='=HYPERLINK("bad")'; self.member.save(); self.client.force_login(self.admin)
        response=self.client.get('/export/'); wb=load_workbook(BytesIO(b''.join(response.streaming_content)))
        self.assertEqual(wb.active['B2'].value,'0012345678'); self.assertEqual(wb.active['A2'].data_type,'s')
    def excel_file(self,rows):
        wb=Workbook(); ws=wb.active
        for row in rows: ws.append(row)
        out=BytesIO(); wb.save(out)
        return SimpleUploadedFile('members.xlsx',out.getvalue())
    def test_excel_preview_commit_and_replay(self):
        self.client.force_login(self.admin)
        rows=[['نام','کد ملی','جواز','واحد','درجه','بدهی'],['عضو تازه','0098765432','002222','واحد تازه','درجه یک',1000]]
        response=self.client.post('/excel/',{'step':'upload','file':self.excel_file(rows)})
        self.assertEqual(response.status_code,200); batch=ImportBatch.objects.get()
        response=self.client.post('/excel/',{'step':'preview','token':batch.token,'map_name':'1','map_national_id':'2','map_number':'3','map_business_name':'4','map_grade':'5','map_opening_balance':'6'})
        self.assertContains(response,'آماده ثبت'); self.assertEqual(Member.objects.count(),1)
        self.client.post('/excel/',{'step':'commit','token':batch.token})
        self.assertEqual(Member.objects.count(),2); self.assertEqual(Due.objects.get().amount,1000)
        self.client.post('/excel/',{'step':'commit','token':batch.token})
        self.assertEqual(Member.objects.count(),2)
    def test_excel_duplicate_rejected(self):
        self.client.force_login(self.admin)
        rows=[['نام','کد','جواز','واحد','درجه'],['تکراری','0012345678','123','آزمایشی','درجه یک']]
        self.client.post('/excel/',{'step':'upload','file':self.excel_file(rows)})
        batch=ImportBatch.objects.get()
        self.client.post('/excel/',{'step':'preview','token':batch.token,'map_name':'1','map_national_id':'2','map_number':'3','map_business_name':'4','map_grade':'5'})
        self.client.post('/excel/',{'step':'commit','token':batch.token})
        self.assertEqual(Member.objects.count(),1)
    def test_document_private_and_invalid_upload_rejected(self):
        self.client.force_login(self.clerk)
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            url=reverse('upload',args=[self.lic.pk])
            self.client.post(url,{'title':'خطرناک','kind':'other','file':SimpleUploadedFile('bad.html',b'<script>alert(1)</script>')})
            self.assertEqual(Document.objects.count(),0)
            self.client.post(url,{'title':'مدرک','kind':'license','file':SimpleUploadedFile('doc.pdf',b'%PDF-1.4\n%%EOF')})
            doc=Document.objects.get(); response=self.client.get(reverse('document',args=[doc.pk])); self.assertEqual(response.status_code,200); response.close()
            self.client.logout(); self.assertEqual(self.client.get(reverse('document',args=[doc.pk])).status_code,302)
    def test_currency_cannot_change_after_financial_data(self):
        self.due(); self.client.force_login(self.admin)
        self.client.post('/settings/',{'action':'settings','title':'اتحادیه','year':1405,'currency':'ریال'})
        self.assertEqual(Settings.current().currency,'تومان')
    def test_payment_retry_is_idempotent(self):
        import uuid
        self.due(amount=100)
        token=uuid.uuid4()
        first=receive(self.accountant,self.lic,100,'card','','1405/01/01',request_id=token)
        second=receive(self.accountant,self.lic,100,'card','','1405/01/01',request_id=token)
        self.assertEqual(first.pk,second.pk); self.assertEqual(Receipt.objects.count(),1)
        self.assertEqual(Allocation.objects.count(),1)
    def test_grade_change_preserves_history_and_old_dues(self):
        other=Grade.objects.create(name='درجه دو'); due=self.due()
        self.client.force_login(self.clerk)
        data={'member-name':self.member.name,'member-national_id':self.member.national_id,'license-number':self.lic.number,'license-business_name':self.lic.business_name,'license-grade':other.pk,'license-active':'on'}
        self.client.post(reverse('member_edit',args=[self.lic.pk]),data)
        self.lic.refresh_from_db(); self.assertEqual(self.lic.grade,other)
        self.assertEqual(GradeHistory.objects.get().grade,other)
        self.assertEqual(due.amount,5000000)
    def test_custom_required_field_enforced_in_excel(self):
        ExtraField.objects.create(key='union_code',label='شناسه اتحادیه',required=True)
        self.client.force_login(self.admin)
        rows=[['نام','کد','جواز','واحد','درجه'],['عضو تازه','0098765432','new','واحد','درجه یک']]
        self.client.post('/excel/',{'step':'upload','file':self.excel_file(rows)})
        batch=ImportBatch.objects.get()
        response=self.client.post('/excel/',{'step':'preview','token':batch.token,'map_name':'1','map_national_id':'2','map_number':'3','map_business_name':'4','map_grade':'5'})
        self.assertContains(response,'ستون‌های الزامی را مشخص کنید.')
        self.assertEqual(Member.objects.count(),1)
    def test_admin_updates_role_without_logging_password(self):
        self.client.force_login(self.admin)
        self.client.post('/users/',{'action':'edit','user':self.clerk.pk,'name':'کاربر آزمایشی','role':'viewer','password':'New-test-password-384!'})
        self.clerk.refresh_from_db(); self.assertTrue(self.clerk.check_password('New-test-password-384!'))
        self.assertTrue(self.clerk.groups.filter(name='viewer').exists())
        self.assertNotIn('New-test-password',str(Audit.objects.first().after))
    def test_personal_report_only_contains_selected_members_licenses(self):
        self.due(amount=1200)
        other=Member.objects.create(name='عضو محرمانه دیگر',national_id='0099999999')
        License.objects.create(member=other,number='OTHER',business_name='واحد دیگر',grade=self.grade)
        License.objects.create(member=self.member,number='SECOND',business_name='واحد دوم همین عضو',grade=self.grade)
        self.client.force_login(self.admin)
        response=self.client.get(reverse('member_report',args=[self.member.pk]))
        self.assertContains(response,self.member.name); self.assertContains(response,'واحد دوم همین عضو')
        self.assertNotContains(response,other.name); self.assertNotContains(response,'OTHER')
        response=self.client.get(reverse('member_report',args=[self.member.pk])+'?format=xlsx')
        book=load_workbook(BytesIO(b''.join(response.streaming_content)))
        self.assertEqual(book.sheetnames,['مشخصات عضو','حق عضویت','دریافت‌ها'])
        self.assertEqual(book.worksheets[0]['B3'].value,'0012345678')
    def test_personal_report_requires_export_permission(self):
        self.client.force_login(self.clerk)
        self.assertEqual(self.client.get(reverse('member_report',args=[self.member.pk])).status_code,403)
    def test_logo_upload_is_validated_and_included_in_report(self):
        from PIL import Image
        buffer=BytesIO(); Image.new('RGB',(24,24),'green').save(buffer,format='PNG')
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            self.client.force_login(self.admin)
            response=self.client.post('/settings/',{'action':'branding','logo':SimpleUploadedFile('logo.png',buffer.getvalue()),'report_address':'نشانی سربرگ','report_phone':'123','report_footer':'گزارش اختصاصی'})
            self.assertEqual(response.status_code,302); self.assertTrue(Settings.current().logo)
            report=self.client.get(reverse('member_report',args=[self.member.pk]))
            self.assertContains(report,'/branding/logo/'); self.assertContains(report,'نشانی سربرگ')
            response=self.client.get(reverse('member_report',args=[self.member.pk])+'?format=xlsx')
            book=load_workbook(BytesIO(b''.join(response.streaming_content)))
            self.assertEqual(len(book.worksheets[0]._images),1)
            self.client.logout(); response=self.client.get('/branding/logo/')
            self.assertEqual(response.status_code,200); response.close()
            self.client.force_login(self.admin)
            old=Settings.current().logo.name
            self.client.post('/settings/',{'action':'branding','logo':SimpleUploadedFile('bad.svg',b'<svg><script>bad</script></svg>')})
            self.assertEqual(Settings.current().logo.name,old)
    def test_branding_changes_admin_only(self):
        self.client.force_login(self.clerk)
        self.assertEqual(self.client.post('/settings/',{'action':'branding','report_footer':'changed'}).status_code,403)
    def test_designer_credit_follows_union_identity_on_login_and_panel(self):
        from PIL import Image
        logo=BytesIO(); Image.new('RGBA',(240,120),(0,100,60,255)).save(logo,format='PNG')
        designer=BytesIO(); Image.new('RGBA',(100,40),(10,30,80,255)).save(designer,format='PNG')
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            self.client.force_login(self.admin)
            self.client.post('/settings/',{'action':'branding','logo':SimpleUploadedFile('union.png',logo.getvalue()),'designer_logo':SimpleUploadedFile('designer.png',designer.getvalue()),'designer_name':'طراح آزمایشی'})
            app=Settings.current(); self.assertTrue(app.designer_logo)
            self.assertEqual(app.designer_name,'طراح آزمایشی')
            self.assertEqual(Audit.objects.first().after['designer_logo'],app.designer_logo.name)
            page=self.client.get('/').content.decode()
            self.assertLess(page.index('class="union-brand"'),page.index('class="designer-credit"'))
            self.assertIn('طراحی و توسعه',page); self.assertIn('طراح آزمایشی',page)
            self.client.logout()
            response=self.client.get('/login/')
            self.assertContains(response,'/branding/designer_logo/'); self.assertContains(response,'طراح آزمایشی')
            response=self.client.get('/branding/designer_logo/'); self.assertEqual(response.status_code,200); response.close()
    def test_designer_is_not_used_as_official_report_logo(self):
        from PIL import Image
        image=BytesIO(); Image.new('RGBA',(60,30),(0,50,100,255)).save(image,format='PNG')
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            self.client.force_login(self.admin)
            self.client.post('/settings/',{'action':'branding','designer_logo':SimpleUploadedFile('designer.png',image.getvalue()),'designer_name':'طراح'})
            response=self.client.get(reverse('member_report',args=[self.member.pk])+'?format=xlsx')
            book=load_workbook(BytesIO(b''.join(response.streaming_content)))
            self.assertEqual(len(book.worksheets[0]._images),0)
            # Updating ordinary settings still serializes its audit with uploaded designer artwork.
            response=self.client.post('/settings/',{'action':'settings','title':'اتحادیه آزمایشی','year':1405,'currency':'تومان'})
            self.assertEqual(response.status_code,302)
            self.assertEqual(Settings.current().title,'اتحادیه آزمایشی')
    def test_unset_designer_credit_is_hidden_and_invalid_art_is_rejected(self):
        self.assertNotContains(self.client.get('/login/'),'class="designer-credit"')
        self.client.force_login(self.admin)
        self.client.post('/settings/',{'action':'branding','designer_logo':SimpleUploadedFile('bad.svg',b'<svg><script>bad</script></svg>'),'designer_name':'نام ثبت‌نشده'})
        self.assertFalse(Settings.current().designer_logo)
        self.assertEqual(Settings.current().designer_name,'')
