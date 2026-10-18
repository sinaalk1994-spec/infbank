import json
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from django.contrib import messages
from django.contrib.auth import login, logout, update_session_auth_hash
from django.contrib.auth.forms import PasswordChangeForm
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User, Group
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError, PermissionDenied
from django.core.paginator import Paginator
from django.db import transaction, IntegrityError
from django.db.models import Q
from django.forms.models import model_to_dict
from django.http import FileResponse, HttpResponse, Http404
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook, load_workbook
from .models import *
from .forms import *
from .services import allowed, require, audit, checked_auth, bill_year, receive, discount, reopen, ROLE_LABELS

def common(request):
    return {'app':Settings.current(),'can_edit':allowed(request.user,'edit'),'can_finance':allowed(request.user,'finance'),'can_discount':allowed(request.user,'discount'),'can_export':allowed(request.user,'export'),'can_audit':allowed(request.user,'audit'),'role_label':'مدیر سامانه' if request.user.is_superuser else next((v for k,v in ROLE_LABELS.items() if request.user.is_authenticated and request.user.groups.filter(name=k).exists()),'کاربر')}

def error_message(exc):
    return '؛ '.join(exc.messages) if isinstance(exc,ValidationError) else 'این رکورد قبلاً ثبت شده یا با اطلاعات موجود تداخل دارد.'

def login_view(request):
    if request.user.is_authenticated: return redirect('dashboard')
    if request.method == 'POST':
        try:
            user=checked_auth(request,request.POST.get('username',''),request.POST.get('password',''))
            login(request,user); audit(user,'ورود به سامانه','session')
            return redirect('dashboard')
        except ValidationError as e: messages.error(request,error_message(e))
    return render(request,'registration/login.html')

@require_POST
def logout_view(request):
    logout(request); return redirect('login')

def get_licenses():
    return License.objects.select_related('member','grade').prefetch_related('dues__allocations','dues__discounts','receipts__allocations').order_by('-id')

def summarize(licenses, year):
    result=[]
    for lic in licenses:
        dues=list(lic.dues.all())
        lic.current_due=next((d for d in dues if d.year==year),None)
        lic.balance=sum(d.balance for d in dues)
        lic.past_balance=sum(d.balance for d in dues if d.year is None or d.year<year)
        lic.credit=sum(r.available for r in lic.receipts.all())
        result.append(lic)
    return result

@login_required
def dashboard(request):
    year=Settings.current().current_year
    all_rows=summarize(get_licenses(),year)
    current=[l.current_due for l in all_rows if l.current_due]
    billed=sum(d.amount for d in current)
    paid=sum(d.paid for d in current)
    settled=sum(d.settled for d in current)
    recent=Receipt.objects.select_related('license__member','actor').order_by('-created_at')[:5]
    return render(request,'registry/dashboard.html',{'nav':'dashboard','title':'نمای کلی','rows':all_rows[:5],'member_count':Member.objects.count(),'license_count':len(all_rows),'total_balance':sum(l.balance for l in all_rows),'current_paid':paid,'settled_count':settled,'progress':round(100*settled/len(current)) if current else 0,'billed':billed,'recent':recent,'recent_audit':Audit.objects.select_related('actor')[:5] if allowed(request.user,'audit') else []})

@login_required
def members(request):
    qs=get_licenses()
    q=request.GET.get('q','').strip()
    if q: qs=qs.filter(Q(member__name__icontains=q)|Q(member__national_id__icontains=q)|Q(number__icontains=q)|Q(business_name__icontains=q))
    grade=request.GET.get('grade','')
    if grade.isdigit(): qs=qs.filter(grade_id=grade)
    rows=summarize(qs,Settings.current().current_year)
    status=request.GET.get('status','')
    if status=='settled': rows=[r for r in rows if r.current_due and r.current_due.settled]
    if status=='unpaid': rows=[r for r in rows if r.current_due and not r.current_due.settled]
    page=Paginator(rows,15).get_page(request.GET.get('page'))
    return render(request,'registry/members.html',{'nav':'members','title':'اعضا و پرونده‌ها','rows':page,'page':page,'count':len(rows),'grades':Grade.objects.all(),'q':q,'status':status,'selected_grade':grade})

@login_required
def member_edit(request,pk=None):
    require(request.user,'edit')
    lic=get_object_or_404(License,pk=pk) if pk else None
    existing=get_object_or_404(Member,pk=request.GET['member']) if not lic and request.GET.get('member','').isdigit() else None
    member=lic.member if lic else existing
    mf=MemberForm(request.POST or None,instance=member,prefix='member')
    lf=LicenseForm(request.POST or None,instance=lic,prefix='license')
    before={'member':model_to_dict(member),'license':model_to_dict(lic)} if lic else {}
    old_grade=lic.grade_id if lic else None
    if request.method=='POST' and mf.is_valid() and lf.is_valid():
        with transaction.atomic():
            obj=lf.save(commit=False); obj.member=mf.save(); obj.save()
            if old_grade != obj.grade_id: GradeHistory.objects.create(license=obj,grade=obj.grade,actor=request.user)
            audit(request.user,'ویرایش پرونده' if pk else 'ایجاد پرونده',f'license:{obj.pk}',before=before,after={'member':model_to_dict(obj.member),'license':model_to_dict(obj)})
        messages.success(request,'پرونده با موفقیت ذخیره شد.'); return redirect('detail',pk=obj.pk)
    return render(request,'registry/member_form.html',{'nav':'members','title':'ویرایش پرونده' if pk else 'پرونده جدید','mf':mf,'lf':lf})

@login_required
def detail(request,pk):
    lic=summarize(get_licenses().filter(pk=pk),Settings.current().current_year)
    if not lic: return get_object_or_404(License,pk=pk)
    lic=lic[0]
    return render(request,'registry/detail.html',{'nav':'members','title':'پرونده عضو','lic':lic,'payment_form':PaymentForm(lic),'debt_form':DebtForm(),'document_form':DocumentForm(),'history':Audit.objects.filter(target=f'license:{pk}').select_related('actor')[:50],'photo':lic.documents.filter(kind='photo').last(),'extra_values':[(f.label,lic.member.extra.get(f.key,'')) for f in ExtraField.objects.all()]})

@login_required
@require_POST
def payment(request,pk):
    require(request.user,'finance'); lic=get_object_or_404(License,pk=pk)
    form=PaymentForm(lic,request.POST)
    if form.is_valid():
        d=form.cleaned_data
        receive(request.user,lic,d['amount'],d['method'],d['reference'],jalali_date(d['paid_on']),d['due'],request_id=d['request_id'])
        messages.success(request,'دریافت ثبت و به بدهی‌ها تخصیص داده شد.')
    else: messages.error(request,str(form.errors.as_text()))
    return redirect('detail',pk=pk)

@login_required
@require_POST
def allocate_credit(request,pk,receipt_id):
    require(request.user,'finance'); lic=get_object_or_404(License,pk=pk)
    receipt=get_object_or_404(Receipt,pk=receipt_id,license=lic)
    receive(request.user,lic,0,'','','',receipt=receipt)
    messages.success(request,'اعتبار آزاد به قدیمی‌ترین بدهی‌ها تخصیص یافت.')
    return redirect('detail',pk=pk)

@login_required
@require_POST
def add_debt(request,pk):
    require(request.user,'finance'); lic=get_object_or_404(License,pk=pk)
    form=DebtForm(request.POST)
    if form.is_valid():
        try:
            with transaction.atomic():
                due=Due.objects.create(license=lic,**form.cleaned_data)
                audit(request.user,'ثبت بدهی / مانده اولیه',f'license:{pk}',after={'due':due.pk,**form.cleaned_data})
            messages.success(request,'بدهی ثبت شد.')
        except IntegrityError: messages.error(request,'برای این سال یا مانده افتتاحیه قبلاً بدهی ثبت شده است.')
    else: messages.error(request,form.errors.as_text())
    return redirect('detail',pk=pk)

@login_required
@require_POST
def add_discount(request,pk,due_id):
    require(request.user,'discount'); due=get_object_or_404(Due,pk=due_id,license_id=pk)
    try:
        value=int(normalize(request.POST.get('value')))
        discount(request.user,due,value,request.POST.get('reason',''),value if request.POST.get('kind')=='percent' else None)
        messages.success(request,'تخفیف تأیید و ثبت شد.')
    except (ValueError,ValidationError) as e: messages.error(request,error_message(e) if isinstance(e,ValidationError) else 'مبلغ یا درصد معتبر وارد کنید.')
    return redirect('detail',pk=pk)

@login_required
@require_POST
def reopen_due(request,pk,due_id):
    require(request.user,'finance'); due=get_object_or_404(Due,pk=due_id,license_id=pk)
    try:
        approver=checked_auth(request,request.POST.get('admin_username',''),request.POST.get('admin_password',''),'approval')
        reopen(request.user,approver,due,request.POST.get('reason',''))
        messages.success(request,'تسویه برگشت خورد. وجوه دریافتی به اعتبار پرونده برگشتند؛ سابقه محفوظ است.')
    except ValidationError as e: messages.error(request,error_message(e))
    return redirect('detail',pk=pk)

@login_required
@require_POST
@transaction.atomic
def upload_document(request,pk):
    require(request.user,'edit'); lic=get_object_or_404(License,pk=pk)
    form=DocumentForm(request.POST,request.FILES)
    if form.is_valid():
        doc=form.save(commit=False); doc.license=lic; doc.uploaded_by=request.user; doc.mime=form.mime; doc.original_name=form.cleaned_data['file'].name[:200]; doc.save()
        audit(request.user,'بارگذاری مدرک',f'license:{pk}',after={'document':doc.pk,'title':doc.title})
        messages.success(request,'مدرک در بایگانی خصوصی ثبت شد.')
    else: messages.error(request,form.errors.as_text())
    return redirect('detail',pk=pk)

@login_required
def document_file(request,pk):
    doc=get_object_or_404(Document,pk=pk)
    audit(request.user,'مشاهده مدرک',f'license:{doc.license_id}',after={'document':pk})
    response=FileResponse(doc.file.open('rb'),content_type=doc.mime,as_attachment=request.GET.get('download')=='1',filename=doc.original_name)
    response['Cache-Control']='private, no-store'; response['Content-Security-Policy']="default-src 'none'; sandbox"
    return response

@login_required
def archive(request):
    docs=Document.objects.select_related('license__member','uploaded_by').order_by('-created_at')
    q=request.GET.get('q','')
    if q: docs=docs.filter(Q(title__icontains=q)|Q(license__member__name__icontains=q)|Q(license__number__icontains=q))
    return render(request,'registry/archive.html',{'nav':'archive','title':'بایگانی مدارک','documents':Paginator(docs,24).get_page(request.GET.get('page')),'q':q})

@login_required
def finance(request):
    rows=summarize(get_licenses(),Settings.current().current_year)
    return render(request,'registry/finance.html',{'nav':'finance','title':'حسابداری و حق عضویت','rows':rows,'receipts':Receipt.objects.select_related('license__member','actor').order_by('-id')[:100],'total_balance':sum(l.balance for l in rows),'total_credit':sum(l.credit for l in rows),'total_received':sum(r.amount for r in Receipt.objects.all()),'total_discount':sum(d.amount for d in Discount.objects.filter(reversed_at=None))})

@login_required
def settings_view(request):
    require(request.user,'admin')
    app=Settings.current()
    if request.method=='POST':
        try:
            with transaction.atomic():
                action=request.POST.get('action')
                if action=='settings':
                    before=model_to_dict(app,exclude=['logo','partner_logo','designer_logo'])
                    year=int(normalize(request.POST.get('year')))
                    if not 1300<=year<=1600: raise ValidationError('سال معتبر نیست.')
                    currency=request.POST.get('currency')
                    if currency not in ['تومان','ریال']: raise ValidationError('واحد پول معتبر نیست.')
                    if currency!=app.currency and (Due.objects.exists() or Receipt.objects.exists() or Tariff.objects.exists()): raise ValidationError('پس از ثبت مبلغ، تغییر واحد پول نیازمند تبدیل حسابرسی‌شده داده‌هاست.')
                    app.title=request.POST.get('title','').strip() or app.title; app.current_year=year; app.currency=currency; app.save()
                    audit(request.user,'تنظیمات سامانه','settings',before,model_to_dict(app,exclude=['logo','partner_logo','designer_logo']))
                elif action=='branding':
                    form=BrandingForm(request.POST,request.FILES,instance=app)
                    before={**model_to_dict(app,exclude=['logo','partner_logo','designer_logo']),'logo':app.logo.name,'partner_logo':app.partner_logo.name,'designer_logo':app.designer_logo.name}
                    if not form.is_valid(): raise ValidationError(form.errors.as_text())
                    app=form.save()
                    audit(request.user,'تغییر هویت و سربرگ','settings',before,{**model_to_dict(app,exclude=['logo','partner_logo','designer_logo']),'logo':app.logo.name,'partner_logo':app.partner_logo.name,'designer_logo':app.designer_logo.name})
                elif action=='grade':
                    name=request.POST.get('name','').strip()
                    if not name: raise ValidationError('نام درجه الزامی است.')
                    Grade.objects.create(name=name)
                    audit(request.user,'تعریف درجه','grade',after={'name':name})
                elif action=='tariff':
                    t=Tariff(year=int(normalize(request.POST.get('year'))),grade_id=int(request.POST.get('grade')),amount=int(normalize(request.POST.get('amount'))))
                    t.full_clean(); t.save(); audit(request.user,'ثبت تعرفه',f'tariff:{t.pk}',after=model_to_dict(t))
                elif action=='field':
                    f=ExtraField(key=request.POST.get('key',''),label=request.POST.get('label',''),required=request.POST.get('required')=='on'); f.full_clean(); f.save()
                    audit(request.user,'تعریف فیلد','field',after=model_to_dict(f))
                elif action=='bill':
                    count=bill_year(request.user,app.current_year)
                    messages.success(request,f'{count} بدهی جدید صادر شد؛ بدهی‌های موجود تغییر نکردند.')
                else: raise ValidationError('عملیات ناشناخته است.')
            messages.success(request,'تغییرات ذخیره شد.')
        except (ValueError,ValidationError,IntegrityError) as e: messages.error(request,error_message(e))
        return redirect('settings')
    return render(request,'registry/settings.html',{'nav':'settings','title':'تنظیمات و تعرفه‌ها','grades':Grade.objects.all(),'tariffs':Tariff.objects.select_related('grade').order_by('-year'),'fields':ExtraField.objects.all(),'branding_form':BrandingForm(instance=app)})

@login_required
def users_view(request):
    require(request.user,'admin')
    if request.method=='POST':
        try:
            with transaction.atomic():
                if request.POST.get('action')=='edit':
                    u=get_object_or_404(User,pk=request.POST.get('user'))
                    if u.is_superuser: raise ValidationError('تغییر مدیر سامانه از این فرم مجاز نیست.')
                    role=request.POST.get('role')
                    if role not in ROLE_LABELS: raise ValidationError('نقش معتبر نیست.')
                    before={'name':u.first_name,'roles':list(u.groups.values_list('name',flat=True))}
                    u.first_name=request.POST.get('name','').strip()[:150]
                    password=request.POST.get('password','')
                    if password: validate_password(password,u); u.set_password(password)
                    u.save(); group,_=Group.objects.get_or_create(name=role); u.groups.set([group])
                    audit(request.user,'ویرایش حساب کاربری',f'user:{u.pk}',before,{'name':u.first_name,'role':role,'password_reset':bool(password)})
                elif request.POST.get('action')=='toggle':
                    u=get_object_or_404(User,pk=request.POST.get('user'))
                    if u.is_superuser: raise ValidationError('غیرفعال‌کردن مدیر از این صفحه مجاز نیست.')
                    u.is_active=not u.is_active; u.save()
                    audit(request.user,'وضعیت کاربر',f'user:{u.pk}',after={'active':u.is_active})
                else:
                    role=request.POST.get('role')
                    if role not in ROLE_LABELS: raise ValidationError('نقش معتبر نیست.')
                    u=User(username=request.POST.get('username','').strip(),first_name=request.POST.get('name','').strip())
                    u.full_clean(exclude=['password']); validate_password(request.POST.get('password',''),u)
                    u.set_password(request.POST['password']); u.save(); group,_=Group.objects.get_or_create(name=role); u.groups.add(group)
                    audit(request.user,'ایجاد کاربر',f'user:{u.pk}',after={'username':u.username,'role':role})
            messages.success(request,'اطلاعات کاربر ذخیره شد.')
        except (ValidationError,IntegrityError) as e: messages.error(request,error_message(e))
        return redirect('users')
    return render(request,'registry/users.html',{'nav':'users','title':'کاربران و دسترسی‌ها','users':User.objects.prefetch_related('groups').all(),'roles':ROLE_LABELS.items()})

@login_required
def audit_view(request):
    require(request.user,'audit')
    entries=Audit.objects.select_related('actor')
    q=request.GET.get('q','')
    if q: entries=entries.filter(Q(action__icontains=q)|Q(actor__username__icontains=q)|Q(target__icontains=q))
    return render(request,'registry/audit.html',{'nav':'audit','title':'ردپای فعالیت‌ها','entries':Paginator(entries,30).get_page(request.GET.get('page')),'q':q})

EXPORT_FIELDS=[('name','نام و نام خانوادگی'),('national_id','کد ملی'),('phone','تلفن همراه'),('number','شماره جواز'),('business_name','نام واحد صنفی'),('grade','درجه'),('year','سال جاری'),('settled','تسویه سال جاری'),('past','بدهی گذشته'),('balance','کل بدهی'),('credit','اعتبار')]
@login_required
def export_excel(request):
    require(request.user,'export')
    selected=request.GET.getlist('field') or [k for k,v in EXPORT_FIELDS]
    fields=[(k,v) for k,v in EXPORT_FIELDS if k in selected]
    wb=Workbook(); ws=wb.active; ws.title='اعضا'; ws.sheet_view.rightToLeft=True
    ws.append([v for k,v in fields])
    year=Settings.current().current_year
    for lic in summarize(get_licenses(),year):
        row={'name':lic.member.name,'national_id':lic.member.national_id,'phone':lic.member.phone,'number':lic.number,'business_name':lic.business_name,'grade':lic.grade.name,'year':year,'settled':'تسویه' if lic.current_due and lic.current_due.settled else 'تسویه نشده' if lic.current_due else 'صادر نشده','past':lic.past_balance,'balance':lic.balance,'credit':lic.credit}
        ws.append([row[k] for k,v in fields])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str): cell.data_type='s'
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=24
    out=BytesIO(); wb.save(out); out.seek(0)
    audit(request.user,'خروجی اکسل','export',after={'fields':selected,'count':ws.max_row-1})
    return FileResponse(out,as_attachment=True,filename='members.xlsx')

IMPORT_FIELDS=[('name','نام و نام خانوادگی',True),('national_id','کد ملی',True),('phone','تلفن همراه',False),('number','شماره جواز',True),('business_name','نام واحد صنفی',True),('grade','نام درجه',True),('opening_balance','مانده بدهی گذشته',False),('opening_year','سال بدهی گذشته (اختیاری)',False),('father_name','نام پدر',False),('birth_date','تاریخ تولد شمسی',False),('home_address','نشانی سکونت',False),('activity','رسته فعالیت',False),('business_address','نشانی واحد صنفی',False),('postal_code','کد پستی',False),('issued_on','تاریخ صدور شمسی',False),('expires_on','تاریخ انقضا شمسی',False)]

@login_required
def import_excel(request):
    require(request.user,'edit')
    import_fields=IMPORT_FIELDS+ [('extra_'+f.key,f.label,f.required) for f in ExtraField.objects.all()]
    context={'nav':'excel','title':'ورود و خروج اکسل','import_fields':import_fields,'export_fields':EXPORT_FIELDS,'mappings':ImportMapping.objects.all()}
    if request.method=='POST':
        try:
            step=request.POST.get('step')
            if step=='upload':
                f=request.FILES.get('file')
                if not f or not f.name.lower().endswith('.xlsx') or f.size>5*1024*1024: raise ValidationError('فایل XLSX حداکثر ۵ مگابایت انتخاب کنید.')
                data=f.read()
                with ZipFile(BytesIO(data)) as z:
                    if sum(i.file_size for i in z.infolist())>40*1024*1024: raise ValidationError('حجم بازشده فایل بیش از حد مجاز است.')
                wb=load_workbook(BytesIO(data),read_only=True,data_only=False)
                sheet=wb.active
                if sheet.max_column>50 or sheet.max_row>1001: raise ValidationError('هر بار حداکثر ۱۰۰۰ ردیف و ۵۰ ستون وارد کنید.')
                raw=[]
                for row in sheet.iter_rows():
                    if any(c.data_type=='f' for c in row): raise ValidationError('فایل ورودی نباید شامل فرمول باشد؛ مقادیر را جایگزین فرمول کنید.')
                    values=[str(c.value) if c.value is not None else '' for c in row]
                    if any(values): raw.append(values)
                wb.close()
                if len(raw)<2: raise ValidationError('فایل باید عنوان ستون‌ها و حداقل یک ردیف داده داشته باشد.')
                batch=ImportBatch.objects.create(actor=request.user,rows=raw)
                context.update(batch=batch,headers=list(enumerate(raw[0])),step='map')
            elif step=='preview':
                batch=get_object_or_404(ImportBatch,token=request.POST.get('token'),actor=request.user,committed_at=None)
                mapping={k:request.POST.get('map_'+k,'') for k,v,r in import_fields}
                if request.POST.get('mapping'):
                    mapping=get_object_or_404(ImportMapping,pk=request.POST['mapping']).mapping
                if any(not mapping.get(k) for k,v,r in import_fields if r): raise ValidationError('ستون‌های الزامی را مشخص کنید.')
                indexes=[v for v in mapping.values() if v]
                if len(indexes)!=len(set(indexes)): raise ValidationError('یک ستون را به چند فیلد متصل نکنید.')
                preview=[]; seen_ids=set(); seen_licenses=set()
                grades={g.name:g.pk for g in Grade.objects.all()}
                for idx,row in enumerate(batch.rows[1:],2):
                    item={k:normalize(row[int(v)-1]) if v and int(v)<=len(row) else '' for k,v in mapping.items()}
                    errors=[]
                    for k,label,required in import_fields:
                        if required and not item.get(k): errors.append(label+' الزامی است')
                    nid=item.get('national_id',''); number=item.get('number','')
                    if not re.fullmatch(r'[0-9]{10}',nid): errors.append('کد ملی باید ۱۰ رقم باشد؛ ستون اکسل را Text کنید')
                    if nid in seen_ids or Member.objects.filter(national_id=nid).exists(): errors.append('عضو تکراری؛ برای افزودن جواز از پرونده موجود استفاده کنید')
                    if number in seen_licenses or License.objects.filter(number=number).exists(): errors.append('شماره جواز تکراری')
                    seen_ids.add(nid); seen_licenses.add(number)
                    if item.get('grade') not in grades: errors.append('درجه تعریف نشده است')
                    limits={'name':150,'number':50,'business_name':150,'phone':20}
                    if any(len(item.get(k,''))>limit for k,limit in limits.items()): errors.append('طول یکی از فیلدها بیش از حد مجاز است')
                    try:
                        amount=int(item.get('opening_balance') or 0); year=int(item['opening_year']) if item.get('opening_year') else None
                        if amount<0 or amount>10**12 or (year and not 1300<=year<Settings.current().current_year): raise ValueError()
                        if year and not amount: raise ValueError()
                        if amount and not allowed(request.user,'finance'): errors.append('ورود بدهی نیازمند دسترسی حسابداری است')
                    except ValueError: errors.append('مبلغ یا سال بدهی گذشته معتبر نیست')
                    member_data={k:item.get(k,'') for k in ['name','national_id','phone','father_name','birth_date']}
                    member_data['address']=item.get('home_address','')
                    member_data.update({k:v for k,v in item.items() if k.startswith('extra_')})
                    license_data={k:item.get(k,'') for k in ['number','business_name','activity','postal_code','issued_on','expires_on']}
                    license_data.update(grade=grades.get(item.get('grade')),address=item.get('business_address',''),active=True)
                    for validation_form in [MemberForm(member_data),LicenseForm(license_data)]:
                        if not validation_form.is_valid():
                            for field,field_errors in validation_form.errors.items():
                                errors.extend(str(e) for e in field_errors)
                    item.update(errors=list(dict.fromkeys(errors)),row=idx); preview.append(item)
                request.session['import_preview']={'token':str(batch.token),'rows':preview}
                if request.POST.get('mapping_name','').strip():
                    ImportMapping.objects.update_or_create(name=request.POST['mapping_name'].strip()[:100],defaults={'mapping':mapping})
                context.update(step='preview',preview=preview,token=batch.token,has_errors=any(r['errors'] for r in preview))
            elif step=='commit':
                saved=request.session.get('import_preview',{})
                if saved.get('token')!=request.POST.get('token') or not saved.get('rows') or any(r['errors'] for r in saved['rows']): raise ValidationError('پیش‌نمایش معتبر موجود نیست.')
                with transaction.atomic():
                    batch=get_object_or_404(ImportBatch.objects.select_for_update(),token=saved['token'],actor=request.user,committed_at=None)
                    for row in saved['rows']:
                        member=Member.objects.create(name=row['name'],national_id=row['national_id'],phone=row.get('phone',''),father_name=row.get('father_name',''),birth_date=jalali_date(row['birth_date']) if row.get('birth_date') else '',address=row.get('home_address',''),extra={f.key:row.get('extra_'+f.key,'') for f in ExtraField.objects.all()})
                        lic=License.objects.create(member=member,number=row['number'],business_name=row['business_name'],grade=Grade.objects.get(name=row['grade']),activity=row.get('activity',''),address=row.get('business_address',''),postal_code=row.get('postal_code',''),issued_on=jalali_date(row['issued_on']) if row.get('issued_on') else '',expires_on=jalali_date(row['expires_on']) if row.get('expires_on') else '')
                        GradeHistory.objects.create(license=lic,grade=lic.grade,actor=request.user)
                        amount=int(row.get('opening_balance') or 0)
                        if amount:
                            require(request.user,'finance')
                            Due.objects.create(license=lic,year=int(row['opening_year']) if row.get('opening_year') else None,amount=amount,basis='مانده اولیه از اکسل؛ دسته '+str(batch.pk))
                        audit(request.user,'ورود پرونده از اکسل',f'license:{lic.pk}',after={'batch':batch.pk,'row':row['row'],'national_id':member.national_id})
                    batch.committed_at=timezone.now(); batch.rows=[]; batch.save()
                request.session.pop('import_preview',None)
                messages.success(request,f"{len(saved['rows'])} پرونده وارد شد."); return redirect('members')
        except (ValidationError,IntegrityError,ValueError,BadZipFile,KeyError,IndexError) as e:
            messages.error(request,error_message(e) if isinstance(e,(ValidationError,IntegrityError)) else 'ساختار فایل یا چیدمان ستون‌ها معتبر نیست.')
    return render(request,'registry/excel.html',context)

@login_required
def password_change(request):
    form=PasswordChangeForm(request.user,request.POST or None)
    if request.method=='POST' and form.is_valid():
        user=form.save(); update_session_auth_hash(request,user)
        audit(user,'تغییر رمز عبور','user:'+str(user.pk))
        messages.success(request,'رمز عبور تغییر کرد.'); return redirect('dashboard')
    return render(request,'registry/password.html',{'title':'تغییر رمز عبور','form':form})


def branding_image(request,kind):
    # Public branding only: never accepts a document ID or arbitrary storage path.
    if kind not in ['logo','partner_logo','designer_logo']: raise Http404
    image=getattr(Settings.current(),kind)
    if not image: raise Http404
    response=FileResponse(image.open('rb'),content_type='image/png')
    response['Cache-Control']='no-cache'
    return response

@login_required
def member_report(request,member_id):
    require(request.user,'export')
    member=get_object_or_404(Member,pk=member_id)
    app=Settings.current()
    licenses=summarize(get_licenses().filter(member=member),app.current_year)
    total_balance=sum(l.balance for l in licenses)
    total_credit=sum(l.credit for l in licenses)
    receipts=Receipt.objects.filter(license__member=member).select_related('license','actor').order_by('paid_on','pk')
    audit(request.user,'خروجی اختصاصی عضو',f'member:{member.pk}',after={'format':'xlsx' if request.GET.get('format')=='xlsx' else 'print','license_ids':[l.pk for l in licenses]})
    if request.GET.get('format')=='xlsx':
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.drawing.image import Image as ExcelImage
        wb=Workbook(); profile=wb.active; profile.title='مشخصات عضو'
        profile.append([app.title]); profile.append(['گزارش اختصاصی عضو',member.name])
        for label,value in [('کد ملی',member.national_id),('تلفن',member.phone),('واحد پول',app.currency),('مانده بدهی',total_balance),('اعتبار آزاد',total_credit),('نشانی اتحادیه',app.report_address),('تلفن اتحادیه',app.report_phone),('توضیحات گزارش',app.report_footer)]: profile.append([label,value])
        streams=[]
        for field,anchor in [(app.logo,'D2'),(app.partner_logo,'F2')]:
            if field:
                with field.open('rb') as source: stream=BytesIO(source.read())
                streams.append(stream); logo=ExcelImage(stream); logo.width=70; logo.height=70; profile.add_image(logo,anchor)
        dues=wb.create_sheet('حق عضویت')
        dues.append(['شماره جواز','واحد صنفی','سال','مبنای بدهی','مبلغ اولیه','تخفیف','پرداخت تخصیص‌یافته','مانده','وضعیت'])
        for lic in licenses:
            for due in lic.dues.all(): dues.append([lic.number,lic.business_name,due.label,due.basis,due.amount,due.discount,due.paid,due.balance,'تسویه' if due.settled else 'باز'])
        payments=wb.create_sheet('دریافت‌ها')
        payments.append(['رسید','جواز','تاریخ','مبلغ','روش','پیگیری','اعتبار آزاد'])
        for r in receipts: payments.append([r.pk,r.license.number,r.paid_on,r.amount,r.get_method_display(),r.reference,r.available])
        for sheet in wb:
            sheet.sheet_view.rightToLeft=True; sheet.freeze_panes='A2'
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value,str): cell.data_type='s'
                    cell.alignment=Alignment(horizontal='right',vertical='center')
            for cell in sheet[1]: cell.fill=PatternFill('solid',fgColor='245A43'); cell.font=Font(color='FFFFFF',bold=True)
            for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=25
            sheet.oddHeader.center.text=app.title.replace('&','&&')
        output=BytesIO(); wb.save(output); output.seek(0)
        response=FileResponse(output,as_attachment=True,filename=f'member-{member.pk}.xlsx')
    else:
        response=render(request,'registry/member_report.html',{'title':'گزارش اختصاصی عضو','member':member,'licenses':licenses,'receipts':receipts,'total_balance':total_balance,'total_credit':total_credit,'generated_at':timezone.localtime()})
    response['Cache-Control']='private, no-store'
    return response
